# Generic utilities script
import pandas as pd
import openpyxl
import os
import socket
from platform import system

# Sharepoint folder, used for storing data and saving outputs
def _get_sharepoint_folder():
    """ Adapted from HCV code, function written by @kelly.maynard"""

    user = socket.gethostname()
    platform = system()

    if user in ['ChrisS-OCT25']:
        folder = r'C:\Users\chris.seaman\Burnet Institute\WG-Modelling-Hepatitis B - Documents\Applications\HepAus Submission Modelling'
    else:
        raise Exception(f'Error: unknown user "{user}", please add user information for future convenience!')

    return os.path.join(os.path.abspath(folder), '')

# GitHub repo, used for accessing frameworks, saving calibrations and databooks.
def _get_github_folder():
    """ Adapted from HCV code, function written by @kelly.maynard"""

    user = socket.gethostname()
    platform = system()

    if user in ['ChrisS-OCT25']:
        folder = r'C:\Users\chris.seaman\Desktop\GitRepos\burnet-hbv'
    else:
        raise Exception(f'Error: unknown user "{user}", please add user information for future convenience!')

    return os.path.join(os.path.abspath(folder), '')

def extract_hbv_effects_by_measure(
        filepath: str,
        sheet_name: str = "Effects - results (Yr1-Yr4)",
) -> dict[str, pd.DataFrame]:
    """
    Parameters
    ----------
    filepath : path to the .xlsx workbook
    sheet_name : name of the effects sheet (Yr1-Yr4)

    Returns
    -------
    dict mapping measure name -> DataFrame (index=year, columns=population)
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))

    measure_row = rows[0]  # measure name, repeated across each 4-year block
    year_row = rows[1]  # year for each column within a block
    data_rows = rows[2:]  # population rows

    # Build column -> (measure, year) map, skipping column 0 (population label)
    col_measure_year = {}
    for col_idx in range(1, len(measure_row)):
        measure = measure_row[col_idx]
        year = year_row[col_idx]
        if measure is None or year is None:
            continue
        col_measure_year[col_idx] = (measure, year)

    # Keep only rows whose population label starts with "HBV"
    hbv_rows = [
        row for row in data_rows
        if row[0] is not None and str(row[0]).strip().upper().startswith("HBV")
    ]

    # Group columns by measure
    measures = {}
    for col_idx, (measure, year) in col_measure_year.items():
        measures.setdefault(measure, []).append((col_idx, year))

    # Keep only measures that are actually HBV measures (their column-block
    # sits under the "HBV ..." header), since HCV measure columns are all
    # None once non-HBV population rows are dropped.
    measures = {
        measure: col_years
        for measure, col_years in measures.items()
        if str(measure).strip().upper().startswith("HBV")
    }

    result = {}
    for measure, col_years in measures.items():
        # sort columns by year for a tidy year-ordered index
        col_years_sorted = sorted(col_years, key=lambda cy: cy[1])
        years = [cy[1] for cy in col_years_sorted]
        cols = [cy[0] for cy in col_years_sorted]

        pop_names = [row[0] for row in hbv_rows]
        data = {
            pop_name: [row[col_idx] for col_idx in cols]
            for pop_name, row in zip(pop_names, hbv_rows)
        }

        df = pd.DataFrame(data, index=years)
        df.index.name = "Year"
        df.columns.name = "Population"
        result[measure] = df

    return result




# Dynamically read and import data from tables within an Excel workbook
def read_table(file_name: str, table_name: str) -> pd.DataFrame:

    """ Allows data to be read from tables in an excel workbook without needing to
    manually specify cell references.

    Taken from: https://stackoverflow.com/questions/54241345/pandas-read-a-table-from-excel"""

    wb = openpyxl.load_workbook(file_name, read_only= False, data_only = True) # openpyxl does not have table info if read_only is True; data_only means any functions will pull the last saved value instead of the formula
    for sheetname in wb.sheetnames: # pulls as strings
        sheet = wb[sheetname] # get the sheet object instead of string
        if table_name in sheet.tables: # tables are stored within sheets, not within the workbook, although table names are unique in a workbook
            tbl = sheet.tables[table_name] # get table object instead of string
            tbl_range = tbl.ref #something like 'C4:F9'
            break # we've got our table, bail from for-loop
    data = sheet[tbl_range] # returns a tuple that contains rows, where each row is a tuple containing cells
    content = [[cell.value for cell in row] for row in data] # loop through those row/cell tuples
    header = content[0] # first row is column headers
    rest = content[1:] # every row that isn't the first is data
    df = pd.DataFrame(rest, columns = header)
    wb.close()
    return df